"""
Generates real sample PDF, DOCX, and XLSX files for verification and testing.
"""
from pathlib import Path
import fitz  # PyMuPDF
import docx
import openpyxl

SAMPLE_DIR = Path(__file__).resolve().parent / "sample_files"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

def generate_pdf():
    pdf_path = SAMPLE_DIR / "enterprise_ai_audit_report.pdf"
    doc = fitz.open()

    # Page 1
    page1 = doc.new_page(width=595, height=842)
    text1 = """Enterprise AI Compliance & Financial Audit Report 2025
Audited Entity: NovaTech Global Systems Inc.
Publication Date: February 18, 2025
Location: Headquartered in Zurich, Switzerland with major operations in New York and Singapore.

1. Executive Summary
NovaTech Global Systems Inc. recorded total annual revenue of $142.5 million in FY 2024, representing an increase of 28.4% compared to FY 2023.
The organization completed the deployment of its Neural Guard cybersecurity architecture across 45 regional datacenters.
CEO Dr. Marcus Vance stated that the company achieved an overall operational efficiency score of 94.6%.

2. Mandatory Security Requirements
All production infrastructure must enforce multi-factor hardware authentication by March 31, 2025.
The data engineering team shall maintain automated snapshot backups every 6 hours with zero data loss tolerance.
Security audits are required to be completed on a quarterly schedule in accordance with ISO 27001 standards."""
    rect1 = fitz.Rect(50, 50, 545, 780)
    page1.insert_textbox(rect1, text1, fontsize=11, fontname="helv")

    # Page 2
    page2 = doc.new_page(width=595, height=842)
    text2 = """3. Capital Allocation and Deadlines
Capital expenditure for Cloud AI infrastructure is budgeted at $38.2 million for Q3 2025.
The submission deadline for the regional compliance filing is strictly no later than November 15, 2025.
The European branch verified that 1,250 enterprise clients renewed multi-year contracts.

4. Audit Findings and Statements
The audit committee concluded that all financial assertions comply with IFRS standards.
The research lab in Berlin developed a proprietary low-latency inference engine.
The engineering teams should ensure that latency remains under 45ms across all European nodes."""
    rect2 = fitz.Rect(50, 50, 545, 780)
    page2.insert_textbox(rect2, text2, fontsize=11, fontname="helv")

    doc.save(str(pdf_path))
    doc.close()
    print(f"Generated PDF: {pdf_path}")

def generate_docx():
    docx_path = SAMPLE_DIR / "system_architecture_requirements.docx"
    doc = docx.Document()

    doc.add_heading("Cloud Data Platform: Architectural Specifications & Requirements", level=1)
    doc.add_paragraph("Document Author: Lead Architect Elena Rostova | Published on January 12, 2025.")
    doc.add_paragraph("The primary datacenter is located in Frankfurt, Germany, with failover capacity located in Amsterdam.")

    doc.add_heading("1. Regulatory and Performance Requirements", level=2)
    doc.add_paragraph("The microservices mesh must sustain 50,000 requests per second with 99.99% availability.")
    doc.add_paragraph("Developers are required to submit vulnerability scan reports before every major release.")
    doc.add_paragraph("The final migration milestone deadline is December 20, 2025, and cannot be extended.")

    doc.add_heading("2. Core Infrastructure Milestones & Budget", level=2)
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Milestone Phase"
    hdr_cells[1].text = "Allocated Budget"
    hdr_cells[2].text = "Target Completion Date"
    hdr_cells[3].text = "Compliance Obligation"

    data = [
        ("Phase 1: Ingestion Engine", "$4,200,000", "April 15, 2025", "Must support Kafka and gRPC"),
        ("Phase 2: Knowledge Graph Integration", "$6,800,000", "August 30, 2025", "Mandatory graph query latency < 15ms"),
        ("Phase 3: Automated Verification Pipeline", "$3,500,000", "November 10, 2025", "Shall achieve 98% claim extraction recall")
    ]

    for item in data:
        row_cells = table.add_row().cells
        for idx, text in enumerate(item):
            row_cells[idx].text = text

    doc.add_paragraph("Overall project budget is capped at $14,500,000.")
    doc.save(str(docx_path))
    print(f"Generated DOCX: {docx_path}")

def generate_xlsx():
    xlsx_path = SAMPLE_DIR / "global_operations_metrics.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Performance_KPIs
    ws1 = wb.active
    ws1.title = "Performance_KPIs"
    ws1.append(["KPI Name", "Target Value", "Actual Q1 Metric", "Variance", "Due Date", "Status"])
    ws1.append(["Customer Acquisition", "15,000 users", "18,450 users", "+23%", "March 31, 2025", "Exceeded"])
    ws1.append(["Server Uptime SLA", "99.95%", "99.98%", "+0.03%", "Ongoing", "Compliant"])
    ws1.append(["Average Resolution Time", "120 minutes", "84 minutes", "-30%", "June 30, 2025", "Met"])
    ws1.append(["Cost Per Transaction", "$0.45", "$0.38", "-15.5%", "December 15, 2025", "Optimized"])

    # Sheet 2: Regional_Offices
    ws2 = wb.create_sheet(title="Regional_Offices")
    ws2.append(["Region", "Headquarters City", "Staff Count", "Annual Budget", "Compliance Lead"])
    ws2.append(["North America", "New York", "450", "$24,000,000", "Sarah Jenkins"])
    ws2.append(["Europe", "London", "320", "$18,500,000", "David Meyer"])
    ws2.append(["Asia Pacific", "Tokyo", "280", "$15,200,000", "Kenji Sato"])

    wb.save(str(xlsx_path))
    print(f"Generated XLSX: {xlsx_path}")

if __name__ == "__main__":
    generate_pdf()
    generate_docx()
    generate_xlsx()
    print("All sample files generated successfully!")
