import uuid
from pathlib import Path
from typing import List, Dict, Any
from .base import BaseParser

class XLSXParser(BaseParser):
    """
    Extracts sheet names, rows, columns, and meaningful cell values from XLSX files using openpyxl.
    """

    def parse(self, file_path: Path, document_id: str) -> List[Dict[str, Any]]:
        import openpyxl

        sections: List[Dict[str, Any]] = []

        try:
            wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
        except Exception as e:
            raise ValueError(f"Unable to read XLSX file. It may be corrupted or invalid: {str(e)}")

        try:
            if not wb.sheetnames:
                raise ValueError("Excel workbook contains no sheets.")

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows_generator = sheet.iter_rows(values_only=True)

                # Find first non-empty row to use as headers
                headers = []
                first_row_num = 1
                for row_idx, row in enumerate(rows_generator, start=1):
                    non_empty = [str(c).strip() if c is not None else "" for c in row]
                    if any(non_empty):
                        headers = non_empty
                        first_row_num = row_idx
                        break

                if not headers:
                    # Sheet is empty
                    continue

                # Now process subsequent rows
                for row_offset, row in enumerate(sheet.iter_rows(min_row=first_row_num + 1, values_only=True), start=first_row_num + 1):
                    row_cells = [str(c).strip() if c is not None else "" for c in row]
                    if not any(row_cells):
                        continue

                    # Create meaningful statements matching header to value
                    cell_pairs = []
                    for col_idx, cell_value in enumerate(row_cells):
                        if not cell_value:
                            continue
                        col_header = headers[col_idx] if col_idx < len(headers) and headers[col_idx] else f"Column {col_idx+1}"
                        cell_pairs.append(f"{col_header}: {cell_value}")

                    if cell_pairs:
                        row_statement = " | ".join(cell_pairs)
                        sections.append({
                            "id": str(uuid.uuid4()),
                            "document_id": document_id,
                            "source": "sheet",
                            "location": f"Sheet '{sheet_name}' (Row {row_offset})",
                            "text": row_statement
                        })

            if not sections:
                raise ValueError("Excel file contains no readable data rows.")

            return sections
        finally:
            wb.close()
