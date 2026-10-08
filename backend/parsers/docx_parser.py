import uuid
from pathlib import Path
from typing import List, Dict, Any
from .base import BaseParser

class DOCXParser(BaseParser):
    """
    Extracts paragraphs and tables from DOCX files using python-docx.
    """

    def parse(self, file_path: Path, document_id: str) -> List[Dict[str, Any]]:
        import docx

        sections: List[Dict[str, Any]] = []

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ValueError(f"Unable to read DOCX file. It may be corrupted or invalid: {str(e)}")

        paragraph_idx = 0
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraph_idx += 1
                sections.append({
                    "id": str(uuid.uuid4()),
                    "document_id": document_id,
                    "source": "paragraph",
                    "location": f"Paragraph {paragraph_idx}",
                    "text": text
                })

        # Extract Tables
        for table_idx, table in enumerate(doc.tables, start=1):
            rows_data = []
            for row in table.rows:
                # Deduplicate cells (merged cells appear multiple times)
                cells = []
                seen_cell_ids = set()
                for cell in row.cells:
                    if id(cell._tc) not in seen_cell_ids:
                        seen_cell_ids.add(id(cell._tc))
                        cells.append(cell.text.strip())
                if any(cells):
                    rows_data.append(cells)

            if not rows_data:
                continue

            # First row may be headers
            headers = rows_data[0]
            for row_idx, row_values in enumerate(rows_data[1:], start=2):
                # Construct formatted row statement: Header: Value | Header: Value
                row_parts = []
                for col_idx, val in enumerate(row_values):
                    if not val:
                        continue
                    header = headers[col_idx] if col_idx < len(headers) and headers[col_idx] else f"Col {col_idx+1}"
                    row_parts.append(f"{header}: {val}")

                row_text = " | ".join(row_parts) if row_parts else " | ".join(row_values)
                if row_text.strip():
                    sections.append({
                        "id": str(uuid.uuid4()),
                        "document_id": document_id,
                        "source": "paragraph",
                        "location": f"Table {table_idx} (Row {row_idx})",
                        "text": row_text
                    })

        if not sections:
            raise ValueError("DOCX document is empty or contains no extractable text.")

        return sections
